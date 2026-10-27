#!/usr/bin/env python3
"""Generate deterministic, legally reusable Office fixtures for local renderer tests.

All generated content is original test data authored for this task (CC0). The
script uses python-docx/openpyxl/python-pptx/xlwt in the test environment only;
it is not an app dependency.
"""
from __future__ import annotations
import io, os, shutil, zipfile, json, struct, random
from pathlib import Path
from PIL import Image, ImageDraw
import docx
from docx.shared import Inches, Pt
import openpyxl
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.util import Inches as PInches
from pptx.enum.text import PP_ALIGN
import xlwt

import argparse

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output', default='fixtures', help='Directory for generated original fixtures.')
_args = parser.parse_args()
OUT = Path(_args.output).resolve()
OUT.mkdir(parents=True, exist_ok=True)

HINDI = 'नमस्ते दुनिया'
UNICODE = 'Café — मस्तिष्क 🧠 — مرحبا'

def image_bytes(color=(204,120,92), size=(320,180), text='Fixture'):
    b=io.BytesIO()
    im=Image.new('RGB',size,color)
    d=ImageDraw.Draw(im)
    d.rectangle((10,10,size[0]-10,size[1]-10),outline=(255,255,255),width=4)
    d.text((24,size[1]//2), text, fill=(255,255,255))
    im.save(b,format='PNG')
    return b.getvalue()

def create_docx(path: Path, rich=False, pages=1):
    d=docx.Document()
    d.core_properties.title='File Claw Office Fixture'
    d.add_heading('File Claw Word fixture',0)
    p=d.add_paragraph('This document verifies the offline DOCX renderer. ')
    p.add_run('Bold text. ').bold=True
    p.add_run('Italic text.').italic=True
    d.add_paragraph(f'Unicode: {HINDI} / {UNICODE}')
    if rich:
        t=d.add_table(rows=3,cols=3)
        for r,row in enumerate(t.rows):
            for c,cell in enumerate(row.cells):
                cell.text=f'R{r+1}C{c+1} {UNICODE if (r+c)%2 else "data"}'
        d.add_heading('Image and structured content',1)
        img=image_bytes(text='DOCX image')
        d.add_picture(io.BytesIO(img),width=Inches(3))
        d.add_paragraph('Bullet one',style='List Bullet')
        d.add_paragraph('Bullet two',style='List Number')
    for i in range(pages):
        d.add_heading(f'Section {i+1}',1)
        d.add_paragraph(('Searchable paragraph with repeatable tokens ALPHA BETA GAMMA. ' * 8) + str(i))
        if i%3==2: d.add_page_break()
    d.save(path)

def create_xlsx(path: Path, rich=False, sheets=1, rows=30, cols=12):
    wb=openpyxl.Workbook()
    ws=wb.active
    ws.title='Summary'
    ws.append(['Name','Value','Formula','Unicode','Notes'])
    for i in range(rows):
        ws.append([f'Item {i+1}',i+1,f'=B{i+2}*2',HINDI if i%7==0 else UNICODE,'ALPHA BETA GAMMA' if i%5==0 else ''])
    if rich:
        ws2=wb.create_sheet('Regional')
        ws2.append(['Region','Sales'])
        for i,n in enumerate(['North','South','East','West']): ws2.append([n,100+i*25])
        chart=BarChart(); chart.title='Sales'; chart.add_data(Reference(ws2,min_col=2,min_row=1,max_row=5),titles_from_data=True); chart.set_categories(Reference(ws2,min_col=1,min_row=2,max_row=5)); ws2.add_chart(chart,'E2')
        ws3=wb.create_sheet('Formula')
        ws3['A1']='=SUM(Summary.B2:B10)'; ws3['A2']='Formula result is shown as cached value where available'
        img=openpyxl.drawing.image.Image(io.BytesIO(image_bytes(color=(80,120,160),text='XLSX image')))
        img.anchor='D2'; ws3.add_image(img)
    for s in range(1,sheets):
        w=wb.create_sheet(f'Sheet{s+1}')
        w['A1']=f'Extra sheet {s+1}'; w['A2']=HINDI
    wb.save(path)

def create_pptx(path: Path, rich=False, slides=3):
    prs=Presentation(); prs.slide_width=PInches(10); prs.slide_height=PInches(5.625)
    for i in range(slides):
        slide=prs.slides.add_slide(prs.slide_layouts[1 if i else 0])
        if i==0:
            slide.shapes.title.text='File Claw Presentation'
            slide.placeholders[1].text=f'Offline renderer fixture • {HINDI}'
        else:
            slide.shapes.title.text=f'Slide {i+1}: searchable content'
            slide.placeholders[1].text='ALPHA BETA GAMMA\n'+UNICODE
        if rich and i==1:
            rows,cols=3,4
            tbl=slide.shapes.add_table(rows,cols,PInches(1),PInches(2.2),PInches(8),PInches(1.6)).table
            for r in range(rows):
                for c in range(cols): tbl.cell(r,c).text=f'T{r+1}C{c+1}'
        if rich and i==2:
            slide.shapes.add_picture(io.BytesIO(image_bytes(color=(80,140,100),text='PPTX image')),PInches(1),PInches(2),width=PInches(4))
            tb=slide.shapes.add_textbox(PInches(5.5),PInches(2),PInches(3),PInches(1.5)); tb.text_frame.text='Embedded image and text\nALPHA BETA'
    prs.save(path)

def variant_from(src: Path, dst: Path, kind: str):
    """Re-label an OOXML package's main content type and add a harmless macro part."""
    with zipfile.ZipFile(src,'r') as zin:
        entries={n:zin.read(n) for n in zin.namelist()}
    ct=entries.get('[Content_Types].xml',b'').decode('utf-8')
    family='word' if 'word/' in ''.join(entries) else ('xl/' if 'xl/' in ''.join(entries) else 'ppt/')
    macros='m' in kind
    templ='t' in kind
    if family=='word':
        main='word/document.xml'; base='application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'; macro='application/vnd.ms-word.document.macroEnabled.main+xml'; template='application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml'
    elif family=='xl/':
        main='xl/workbook.xml'; base='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'; macro='application/vnd.ms-excel.sheet.macroEnabled.main+xml'; template='application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml'
    else:
        main='ppt/presentation.xml'; base='application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml'; macro='application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml'; template='application/vnd.openxmlformats-officedocument.presentationml.template.main+xml'
    replacement=macro if macros else (template if templ else base)
    ct=ct.replace(base,replacement)
    if macros:
        ct=ct.replace('</Types>','<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
        macro_name={'word':'word/vbaProject.bin','xl/':'xl/vbaProject.bin'}.get(family,'ppt/vbaProject.bin')
        entries[macro_name]=b'NOT_EXECUTABLE_TEST_MACRO'
    # Persist the rewritten content types; the variants are not real Word/Excel/
    # PowerPoint packages without this step.
    entries['[Content_Types].xml']=ct.encode('utf-8')
    with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as zout:
        for name,data in entries.items(): zout.writestr(name,data)

def create_adversarial(src: Path):
    # truncated, empty, misleading extension
    data=src.read_bytes()
    (OUT/'truncated.docx').write_bytes(data[:max(64,len(data)//3)])
    (OUT/'empty.xlsx').write_bytes(b'')
    shutil.copyfile(src,OUT/'misleading.pdf')  # OOXML bytes with PDF extension
    # malformed/corrupt zip
    (OUT/'corrupt.pptx').write_bytes(b'PK\x03\x04corrupted-not-a-real-zip')
    # path traversal and XXE in a plausible Word package
    with zipfile.ZipFile(src,'r') as zin: entries={n:zin.read(n) for n in zin.namelist()}
    entries['../../evil.txt']=b'path traversal'
    entries['word/document.xml']=b'''<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&xxe; XXE ALPHA</w:t></w:r></w:p></w:body></w:document>'''
    with zipfile.ZipFile(OUT/'adversarial.docx','w',zipfile.ZIP_DEFLATED) as z:
        for n,b in entries.items(): z.writestr(n,b)
    # ZIP-bomb-like entry (50 MiB of zeros, small compressed size)
    bomb=OUT/'zip-bomb.docx'
    with zipfile.ZipFile(bomb,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        z.writestr('[Content_Types].xml',b'<Types/>')
        z.writestr('word/document.xml',b'<w:document/>')
        z.writestr('word/media/huge.bin',b'\0'*(50*1024*1024))
    # external relationship
    with zipfile.ZipFile(src,'r') as zin: entries={n:zin.read(n) for n in zin.namelist()}
    rels='word/_rels/document.xml.rels' if 'word/document.xml' in entries else ('xl/_rels/workbook.xml.rels' if 'xl/workbook.xml' in entries else 'ppt/_rels/presentation.xml.rels')
    entries[rels]=b'''<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdExt" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/payload" TargetMode="External"/></Relationships>'''
    with zipfile.ZipFile(OUT/'external-rel.docx','w',zipfile.ZIP_DEFLATED) as z:
        for n,b in entries.items(): z.writestr(n,b)

create_docx(OUT/'basic.docx')
create_docx(OUT/'rich.docx',rich=True,pages=12)
create_xlsx(OUT/'basic.xlsx')
create_xlsx(OUT/'rich.xlsx',rich=True,sheets=3,rows=120,cols=20)
create_xlsx(OUT/'large.xlsx',rows=5000,cols=40)
create_pptx(OUT/'basic.pptx',slides=2)
create_pptx(OUT/'rich.pptx',rich=True,slides=8)
for src,kind in [('basic.docx','m'),('basic.docx','t'),('basic.xlsx','m'),('basic.xlsx','t'),('basic.pptx','m'),('basic.pptx','t')]:
    suffix={'m':'m','t':'x'}[kind]
    stem=OUT/src
    # docm/dotx, xlsm/xltx, pptm/potx
    base=stem.stem
    ext={'docx':('docm','dotx'),'xlsx':('xlsm','xltx'),'pptx':('pptm','potx')}[stem.suffix[1:]]
    dst=OUT/f'{base}.{ext[0] if kind=="m" else ext[1]}'
    variant_from(stem,dst,kind)
# legacy XLS via xlwt
wb=xlwt.Workbook(); sh=wb.add_sheet('Legacy'); sh.write(0,0,'Name'); sh.write(0,1,'Value'); sh.write(1,0,HINDI); sh.write(1,1,42); wb.save(OUT/'legacy.xls')
create_adversarial(OUT/'basic.docx')
print(json.dumps({'root':str(OUT),'files':sorted([str(p.relative_to(OUT)) for p in OUT.rglob('*') if p.is_file()]),'sizes':{p.name:p.stat().st_size for p in OUT.iterdir() if p.is_file()}},ensure_ascii=False,indent=2))
